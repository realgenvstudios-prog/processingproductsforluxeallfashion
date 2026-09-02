"""Exports the product table to a review-ready spreadsheet, with every image
from the post's carousel embedded (not just the first one)."""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from db import get_conn

OUT_PATH = Path(__file__).resolve().parents[1] / "data" / "catalog_review.xlsx"
THUMB_DIR = Path(__file__).resolve().parents[1] / "data" / "thumbs"
THUMB_WIDTH = 120
ROW_HEIGHT = 90

DATA_COLUMNS = [
    "product_id", "is_product_post", "product_name", "brand", "category",
    "description", "price", "currency", "sizes", "colors", "availability_status",
    "confidence_product_name", "confidence_brand", "confidence_price", "confidence_availability",
    "review_required", "review_reason", "duplicate_of", "dedup_score",
    "instagram_url", "caption", "post_date", "media_type",
]


def make_thumbnail(source_path: str, dest_path: Path):
    with PILImage.open(source_path) as im:
        im = im.convert("RGB")
        ratio = THUMB_WIDTH / im.width
        im = im.resize((THUMB_WIDTH, int(im.height * ratio)))
        im.save(dest_path, "JPEG", quality=80)


def export():
    conn = get_conn()
    rows = conn.execute(
        """SELECT pr.*, p.post_url, p.caption, p.post_date, p.media_type
           FROM product pr
           JOIN instagram_post p ON p.id = pr.post_id
           ORDER BY pr.id"""
    ).fetchall()

    max_media = conn.execute(
        "SELECT MAX(cnt) FROM (SELECT post_id, COUNT(*) as cnt FROM media GROUP BY post_id)"
    ).fetchone()[0] or 1
    image_columns = [f"image_{i+1}" for i in range(max_media)]

    THUMB_DIR.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "catalog"
    ws.append(image_columns + DATA_COLUMNS)
    for col_idx in range(1, len(image_columns) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 18
    for col_idx in range(len(image_columns) + 1, len(image_columns) + len(DATA_COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 18

    for i, r in enumerate(rows):
        row_num = i + 2
        confidence = json.loads(r["confidence_json"]) if r["confidence_json"] else {}
        record = {
            "product_id": r["id"],
            "is_product_post": bool(r["is_product_post"]),
            "product_name": r["product_name"],
            "brand": r["brand"],
            "category": r["category"],
            "description": r["description"],
            "price": r["price"],
            "currency": r["currency"],
            "sizes": r["sizes"],
            "colors": r["colors"],
            "availability_status": r["availability_status"],
            "confidence_product_name": confidence.get("product_name"),
            "confidence_brand": confidence.get("brand"),
            "confidence_price": confidence.get("price"),
            "confidence_availability": confidence.get("availability_status"),
            "review_required": bool(r["review_required"]),
            "review_reason": r["review_reason"],
            "duplicate_of": r["duplicate_of"],
            "dedup_score": r["dedup_score"],
            "instagram_url": r["post_url"],
            "caption": r["caption"],
            "post_date": r["post_date"],
            "media_type": r["media_type"],
        }
        for col_idx, col_name in enumerate(DATA_COLUMNS):
            ws.cell(row=row_num, column=len(image_columns) + col_idx + 1, value=record.get(col_name))

        media_rows = conn.execute(
            "SELECT local_path FROM media WHERE post_id = ? ORDER BY position",
            (r["post_id"],),
        ).fetchall()
        ws.row_dimensions[row_num].height = ROW_HEIGHT

        for m_idx, m in enumerate(media_rows):
            col_letter = get_column_letter(m_idx + 1)
            thumb_path = THUMB_DIR / f"{r['id']}_{m_idx}.jpg"
            try:
                make_thumbnail(m["local_path"], thumb_path)
                img = XLImage(str(thumb_path))
                ws.add_image(img, f"{col_letter}{row_num}")
            except Exception as e:
                ws.cell(row=row_num, column=m_idx + 1, value=f"[image failed: {e}]")

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_PATH)
    conn.close()
    print(f"exported {len(rows)} row(s), up to {max_media} image(s) each, to {OUT_PATH}")


if __name__ == "__main__":
    export()
